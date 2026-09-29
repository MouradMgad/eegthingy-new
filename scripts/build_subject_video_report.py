from __future__ import annotations

import json
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "paper results"
THUMBS = ROOT / "report_work" / "thumbs"
OUTPUT = RESULTS / "Subject_6_and_10_Live_Maze_Report.docx"

ACCENT = "A61B1B"
HEADER = "3B3B3B"
BORDER = "D9D9D9"
PALE = "F5F5F5"
MUTED = RGBColor(89, 89, 89)

SESSION_FOLDERS = [
    "subject_10_test_1",
    "subject_10_test_2",
    "subject_6_test_1",
    "subject_6_test_2",
]


def load_sessions() -> list[dict]:
    sessions = []
    for folder_name in SESSION_FOLDERS:
        folder = RESULTS / folder_name
        data = json.loads((folder / "results summary.json").read_text(encoding="utf-8"))
        session = data["session"]
        wrong = [
            int(row["turn"])
            for row in data["corners"]
            if row.get("scored") is True and row.get("result") == "wrong"
        ]
        sessions.append(
            {
                "folder_name": folder_name,
                "folder": folder,
                "subject": int(session["subject"]),
                "test": int(session["test"]),
                "maze": int(session["maze_number"]),
                "seed": int(session["maze_seed"]),
                "correct": int(session["correct"]),
                "wrong": int(session["wrong"]),
                "accuracy": float(session["accuracy"]),
                "recovery": int(session["recovery_attempts"]),
                "duration_s": float(session["wall_duration_s"]),
                "start": datetime.fromisoformat(session["start_time"]),
                "end": datetime.fromisoformat(session["end_time"]),
                "checkpoint": Path(session["checkpoint"]).name,
                "wrong_corners": wrong,
                "face_video": folder / "camera.mp4",
                "maze_video": folder / "maze_session_replay.mp4",
                "face_image": THUMBS / f"{folder_name}_face.jpg",
                "maze_image": THUMBS / f"{folder_name}_maze.jpg",
            }
        )
    return sessions


def fmt_duration(seconds: float) -> str:
    total = int(round(seconds))
    return f"{total // 60}:{total % 60:02d}"


def fmt_percent(correct: int, total: int) -> str:
    value = (Decimal(correct) * Decimal(100) / Decimal(total)).quantize(
        Decimal("0.1"), rounding=ROUND_HALF_UP
    )
    return f"{value}%"


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for existing in list(p_pr.findall(qn("w:pBdr"))):
        p_pr.remove(existing)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    p_pr.append(borders)


def set_font(run, name: str = "Aptos", size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, target: Path, *, bold=False, size=9.5) -> None:
    r_id = paragraph.part.relate_to(target.resolve().as_uri(), RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    hyperlink.set(qn("w:tooltip"), str(target))
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    r_pr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_hyperlinked_picture(paragraph, image: Path, video: Path, width: float, alt: str) -> None:
    run = paragraph.add_run()
    run.add_picture(str(image), width=Inches(width))
    for node in run._r.xpath(".//wp:docPr"):
        node.set("descr", alt)
        node.set("title", alt)
    r_id = paragraph.part.relate_to(video.resolve().as_uri(), RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    hyperlink.set(qn("w:tooltip"), str(video))
    run._r.getparent().remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.56)
    section.left_margin = Inches(0.67)
    section.right_margin = Inches(0.67)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(31, 31, 31)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(5)

    for style_name, size, before, after in (
        ("Heading 1", 18, 0, 7),
        ("Heading 2", 12.5, 8, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    run = paragraph.add_run("Live Maze Evaluation   |   26 September 2026   |   ")
    set_font(run, size=8)
    run.font.color.rgb = MUTED
    page_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    page_run._r.extend([begin, instr, separate, value, end])
    set_font(page_run, size=8)
    page_run.font.color.rgb = MUTED


def add_label_paragraph(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(label)
    set_font(run, size=10.5)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    body = paragraph.add_run(text)
    set_font(body, size=10.5)


def add_results_table(doc: Document, sessions: list[dict]) -> None:
    headers = ["Session", "Maze seed", "Duration", "Correct", "Accuracy", "Wrong corners", "Recovery moves"]
    widths = [0.78, 0.78, 0.72, 0.72, 0.72, 2.22, 0.88]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, HEADER)
        set_cell_margins(cell, top=90, bottom=90, start=70, end=70)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            set_font(run, size=8.5)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, item in enumerate(sessions, start=1):
        wrong_text = ", ".join(str(x) for x in item["wrong_corners"]) or "None"
        values = [
            f"S{item['subject']} T{item['test']}",
            f"M{item['maze']}  S{item['seed']}",
            fmt_duration(item["duration_s"]),
            f"{item['correct']}/40",
            f"{item['accuracy'] * 100:.1f}%",
            wrong_text,
            str(item["recovery"]),
        ]
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = value
            set_cell_shading(cell, PALE if row_index % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=92, bottom=92, start=70, end=70)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 5 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_font(run, size=8.5)
                if index == 4:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(ACCENT)
    set_table_widths(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)


def add_video_pair(doc: Document, item: dict) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [3.45, 3.45]
    for index, (kind, image, video, duration) in enumerate(
        (
            ("Face camera", item["face_image"], item["face_video"], None),
            ("Maze replay", item["maze_image"], item["maze_video"], None),
        )
    ):
        cell = table.cell(0, index)
        set_cell_borders(cell, color="FFFFFF", size="0")
        set_cell_margins(cell, top=0, start=35, bottom=0, end=35)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        add_hyperlinked_picture(
            p,
            image,
            video,
            width=3.30,
            alt=f"Subject {item['subject']} test {item['test']} {kind.lower()} video thumbnail",
        )
        label = cell.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label.paragraph_format.space_after = Pt(0)
        label.paragraph_format.keep_with_next = True
        run = label.add_run(kind)
        set_font(run, size=9.5)
        run.bold = True
        label.add_run("  ")
        link_text = "Open original video" if kind == "Face camera" else "Open reconstructed video"
        add_hyperlink(label, link_text, video, bold=False, size=9.5)
    set_table_widths(table, widths)


def add_subject_page(doc: Document, subject: int, sessions: list[dict]) -> None:
    doc.add_page_break()
    heading = doc.add_heading(f"Subject {subject}", level=1)
    heading.paragraph_format.space_after = Pt(4)
    total_correct = sum(item["correct"] for item in sessions)
    total_trials = 40 * len(sessions)
    accuracy = fmt_percent(total_correct, total_trials)
    checkpoint = sessions[0]["checkpoint"]
    descriptor = "pretrained seven subject checkpoint" if subject == 10 else "subject session fine tuned checkpoint"
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.add_run(
        f"Across two completed mazes, Subject {subject} scored {total_correct}/{total_trials} "
        f"({accuracy}) using the {descriptor} {checkpoint}."
    )
    for run in p.runs:
        set_font(run, size=10.25)

    for index, item in enumerate(sessions):
        heading = doc.add_heading(f"Test {item['test']} Maze {item['maze']} Seed {item['seed']}", level=2)
        heading.paragraph_format.page_break_before = False
        heading.paragraph_format.keep_with_next = True
        wrong_text = ", ".join(str(x) for x in item["wrong_corners"]) or "none"
        stats = doc.add_paragraph()
        stats.paragraph_format.space_after = Pt(4)
        stats.paragraph_format.keep_with_next = True
        stats.add_run(
            f"{item['correct']}/40 correct  |  {item['accuracy'] * 100:.1f}% accuracy  |  "
            f"wrong corners {wrong_text}  |  {item['recovery']} recovery moves  |  "
            f"session {fmt_duration(item['duration_s'])}"
        )
        for run in stats.runs:
            set_font(run, size=9.25)
            run.font.color.rgb = MUTED
        add_video_pair(doc, item)
        if index == 0:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(1)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(0)
    note_run = note.add_run(
        "Click a thumbnail or its link to open the corresponding MP4 file on this computer."
    )
    set_font(note_run, size=8.5)
    note_run.italic = True
    note_run.font.color.rgb = MUTED


def build() -> Path:
    sessions = load_sessions()
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.add_run("Live Maze Evaluation Report")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(13)
    run = subtitle.add_run("MIRepNet results for Subjects 10 and 6")
    set_font(run, size=12)
    run.font.color.rgb = MUTED

    total_correct = sum(item["correct"] for item in sessions)
    total_trials = 40 * len(sessions)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    intro.add_run(
        "All four live MIRepNet maze tests finished successfully. Using the first model decision "
        "at each planned corner, Subject 10 achieved 82.5% accuracy (66/80), Subject 6 achieved "
        f"96.3% (77/80), and combined accuracy was {100 * total_correct / total_trials:.1f}% "
        f"({total_correct}/{total_trials})."
    )

    doc.add_heading("Results Overview", level=1)
    add_results_table(doc, sessions)

    doc.add_heading("Scoring and Video Record", level=1)
    add_label_paragraph(
        doc,
        "Scoring  ",
        "Each planned corner contributes one result. Extra physical turns used to recover from an "
        "incorrect decision are reported as recovery moves and are excluded from accuracy.",
    )
    add_label_paragraph(
        doc,
        "Videos  ",
        "The face clips are the original camera recordings. The maze clips are live style replays "
        "reconstructed from the recorded decisions and timing after the original maze recordings "
        "were corrupted. They reproduce the path, model choices, turn numbers and on screen status "
        "messages, but they are not the original captured maze frames.",
    )

    doc.add_heading("Subject Comparison", level=1)
    comparison = doc.add_paragraph()
    comparison.paragraph_format.space_after = Pt(0)
    comparison.add_run(
        "Subject 6 recorded 11 more correct corners than Subject 10 across the same number of scored "
        "corners (77 versus 66). The sessions used different model checkpoints and different mazes, "
        "so this descriptive comparison does not separate participant, checkpoint or maze effects."
    )

    add_subject_page(doc, 10, [item for item in sessions if item["subject"] == 10])
    add_subject_page(doc, 6, [item for item in sessions if item["subject"] == 6])

    core = doc.core_properties
    core.title = "Live Maze Evaluation Report"
    core.subject = "MIRepNet live maze results for Subjects 10 and 6"
    core.author = "MIRepNet Research Team"
    core.keywords = "MIRepNet, EEG, BCI, motor imagery, live maze, Subject 10, Subject 6"
    core.comments = "Includes linked face-camera recordings and reconstructed maze replay videos."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    build()
