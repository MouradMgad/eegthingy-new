"""Build the MIRepNet methods and implementation disclosure for the paper."""

from __future__ import annotations

import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper results" / "MIRepNet_IEEE_Implementation_and_Methods.docx"
BLACK = RGBColor(0, 0, 0)
RED = RGBColor(157, 24, 31)
GRAY = RGBColor(84, 84, 84)
LIGHT = "F5F5F5"
BORDER = "D9D9D9"


def style(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.58), Inches(0.54)
    sec.left_margin, sec.right_margin = Inches(0.72), Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Aptos", Pt(10), RGBColor(25, 25, 25)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (
        ("Title", 23, 0, 5),
        ("Heading 1", 15, 9, 4),
        ("Heading 2", 11.5, 7, 3),
    ):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.color.rgb = "Aptos Display", Pt(size), BLACK
        st.font.bold = name != "Title"
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("MIRepNet implementation methods  |  26 September 2026  |  ").font.size = Pt(8)
    r = footer.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.text = " PAGE "
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t"); value.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.extend([begin, instr, sep, value, end])
    r.font.size = Pt(8)
    for run in footer.runs:
        run.font.color.rgb = GRAY


def no_title_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for old in list(ppr.findall(qn("w:pBdr"))):
        ppr.remove(old)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    ppr.append(borders)


def para(doc: Document, text: str = "", *, small=False, keep=False) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.keep_together = keep
    if small:
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = GRAY


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.add_run(text)


def heading(doc: Document, text: str, level=1) -> None:
    doc.add_heading(text, level=level)


def equation(doc: Document, expression: str, label: str) -> None:
    """Insert a native OOXML Office Math paragraph, not a raster equation."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    omath_para = OxmlElement("m:oMathPara")
    omath = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = expression
    run.append(text)
    omath.append(run)
    omath_para.append(omath)
    p._p.append(omath_para)
    caption = doc.add_paragraph(label)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(5)
    for r in caption.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = GRAY


def cell_format(cell, header=False, shade=False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "85" if edge in ("top", "bottom") else "105")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)
        borders.append(node)
    tcpr.append(borders)
    if header or shade:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "414141" if header else LIGHT)
        tcpr.append(shd)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.size = Pt(8.5)
            if header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
        cell_format(t.rows[0].cells[i], header=True)
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))
    for j, row in enumerate(rows):
        cells = t.add_row().cells
        for i, item in enumerate(row):
            cells[i].text = item
            cell_format(cells[i], shade=bool(j % 2))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def percent(n: int, d: int) -> str:
    return str((Decimal(n) * 100 / d).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)) + "%"


def session_scores() -> list[dict]:
    out = []
    for folder in ("subject_10_test_1", "subject_10_test_2", "subject_6_test_1", "subject_6_test_2"):
        report = json.loads((ROOT / "paper results" / folder / "results summary.json").read_text())
        s = report["session"]
        out.append({"subject": int(s["subject"]), "test": int(s["test"]),
                    "maze": int(s["maze_number"]), "seed": int(s["maze_seed"]),
                    "correct": int(s["correct"]), "wrong": int(s["wrong"]),
                    "recovery": int(s["recovery_attempts"]), "duration": float(s["wall_duration_s"])})
    return out


def main() -> None:
    doc = Document()
    style(doc)
    p = doc.add_paragraph(style="Title")
    no_title_border(p)
    p.add_run("MIRepNet Model Implementation and Methods")
    sub = doc.add_paragraph()
    sub.add_run("Technical disclosure for the EEG motor imagery live maze study").font.color.rgb = GRAY
    para(doc, "Implementation snapshot: 26 September 2026. This document describes the active classifier/mirepnet.py path and the checkpoints recorded for Subjects 10 and 6. The upstream MIRepNet foundation weights were released by Liu et al. [1] and distributed in Braindecode format [2]. The project fine tuned those weights on local recordings; it did not reproduce upstream pretraining.")

    heading(doc, "Study and Implementation Scope")
    para(doc, "The task is binary left versus right hand motor imagery. Class 1 means left hand imagery and class 2 means right hand imagery; the live maze maps these to left and right turns. The system streams OpenBCI Cyton plus Daisy EEG, preprocesses a task window, applies MIRepNet, recenters its decision score using unlabeled calibration, commits a temporally stable command, and scores one model decision per planned maze corner. Four fixed mazes have 40 scored corners each.")
    para(doc, "The operational model is MIRepNetModel in classifier/mirepnet.py. An older alternative under mirepnet_pipeline uses a learnable 15 to 45 channel adapter. That adapter is absent from the two live checkpoints reported here; the active path uses deterministic inverse distance interpolation. All numerical results below refer to the active path.")

    heading(doc, "EEG Acquisition and Trial Definition")
    table(doc, ["Item", "Implementation used"], [
        ["Hardware and rate", "OpenBCI Cyton plus Daisy; 16 streamed EEG channels at 125 Hz in the four recorded sessions."],
        ["Observed montage", "15 selected channels: Cz, Pz, C3, C4, T5, T6, Fz, F7, F8, F3, F4, T3, T4, P3, P4. Cyton channel 8 is marked DEAD and excluded."],
        ["Aliases", "T3→T7, T4→T8, T5→P7, T6→P8, matching the model coordinate/template convention."],
        ["Offline labels", "Annotated left_hand/task→1 and right_hand/task→2 from local Exp4 FIF recordings."],
        ["Band and rate", "8–30 Hz zero phase FIR filter using MNE; resample to 250 Hz."],
        ["Epoch", "Task interval 0–2 s yields 500 samples at 250 Hz; this same segment is concatenated with itself to make 1,000 samples (nominal 4 s)."],
        ["Center and reject", "Subtract each channel's temporal mean; offline reject a trial when any channel has peak to peak amplitude ≥100 μV. Live windows are not rejected by this offline rule."],
    ], [1.42, 5.62])
    para(doc, "The repeated input contains two copies of the same two second task EEG. Thus 1,000 samples meet the pretrained tensor shape but do not provide four seconds of independent neural evidence. The alternative plan task mode includes the two seconds before the task event; it was not the default for these checkpoints.")

    heading(doc, "Spatial Alignment and Channel Projection")
    para(doc, "Euclidean Alignment is fitted in the 15 observed channels before interpolation. For trial i, the implementation computes a temporal mean per channel and sample covariance Cᵢ. It averages Cᵢ across calibration trials, clips small eigenvalues, and multiplies each trial by the inverse square root of that reference. In training, each recording receives its own alignment reference. A global training reference is also saved in the checkpoint for later inference.")
    equation(doc, "Cᵢ = (Xᵢ − μᵢ)(Xᵢ − μᵢ)ᵀ / (T − 1)    C̄ = (1/N) ∑ᵢ Cᵢ", "(1) Trial covariance and mean reference covariance")
    equation(doc, "C̄ = VΛVᵀ    R = V diag(max(λⱼ, ε)⁻¹ᐟ²)Vᵀ    Xᵢᴱᴬ = RXᵢ", "(2) Alignment; ε = max(10⁻⁷ λₘₐₓ, 10⁻¹⁵)")
    para(doc, "The 45 by 15 projection matrix W retains an exact channel if measured. An absent template channel is a weighted average of measured channels using the published two dimensional coordinates. This maps 15 aligned signals into the fixed 45 channel order. The output rank cannot exceed 15, so interpolation does not create new independent electrodes.")
    equation(doc, "Wⱼₖ = dⱼₖ⁻¹ / ∑ₗ dⱼₗ⁻¹    dⱼₖ = ‖pⱼ − pₖ‖₂ + 10⁻⁶    X⁴⁵ = WXᴱᴬ", "(3) Inverse distance interpolation for a template channel without an exact match")

    heading(doc, "Network Architecture")
    para(doc, "The active model has 5,142,786 trainable parameters. Its input is a batch of 45 × 1,000 EEG values. The convolutional front end creates 61 temporal tokens of width 256; six transformer blocks encode them; mean pooling and a two class linear head produce logits.")
    table(doc, ["Stage", "Operation and output", "Trainable parameters"], [
        ["Temporal convolution", "Conv2d 1→64, kernel (1,25), no padding; 45 × 976 feature map", "1,664"],
        ["Spatial convolution", "Conv2d 64→128, kernel (45,1); 1 × 976 feature map", "368,768"],
        ["Normalization and tokens", "BatchNorm2d(128), ELU, average pool (1,75) stride (1,15), dropout 0.5, 1×1 projection 128→256; 61 tokens", "33,280"],
        ["Transformer encoder", "Six pre normalization residual blocks; 8 attention heads of width 32; feed forward 256→1,024→256 with GELU; dropout 0.5", "4,738,560"],
        ["Classification", "Mean over 61 tokens; linear 256→2 logits", "514"],
    ], [1.34, 4.58, 1.12])
    para(doc, "The embedding subtotal is 403,712 parameters. The implementation has no explicit learned positional embedding in its forward pass. Each attention block applies layer normalization before attention and before the feed forward network, with a residual connection around each branch. Attention divides QKᵀ by √256 (the full embedding width), as in this code, rather than by √32.")
    equation(doc, "A = softmax(QKᵀ / √256)V    H′ = H + Dropout(Proj(A))", "(4) Multihead self attention branch, with eight heads concatenated before Proj")
    equation(doc, "h = (1/61) ∑ₜ Hₜ⁽⁶⁾    ℓ = Wₒh + bₒ    pₖ = exp(ℓₖ) / ∑ⱼ exp(ℓⱼ)", "(5) Mean pooled representation, binary logits and softmax probabilities")

    heading(doc, "Weight Provenance and Local Fine Tuning")
    para(doc, "The upstream model was pretrained by its original authors with supervised motor imagery classification and masked token reconstruction [1]. The project loads a pinned Braindecode safetensors rehost (revision 857f1e3642976be9f2ca6883e5508c3f7c91d86f; SHA 256 a3a7eb0f100b703cd40ee2ed634dbc04d1f6ae04931404299dc2f81c4175559b) [2]. It loads the convolutional embedding and transformer weights, discards the released three class output head, and initializes a new two class head. No reconstruction decoder or masked token objective runs during local fine tuning or live inference.")
    table(doc, ["Setting", "Active value or rule"], [
        ["Objective", "Two class cross entropy; no mixup, augmentation, label smoothing, or class weighting in this active training path."],
        ["Optimizer", "AdamW; encoder learning rate 1×10⁻⁴; head learning rate 1×10⁻³; weight decay 1×10⁻⁴."],
        ["Schedule", "CosineAnnealingLR stepped after each epoch; T_max=20; maximum 20 epochs; batch size 16."],
        ["Freeze schedule", "Embedding and all transformer weights frozen for epochs 1–2; the new head trains. Starting epoch 3, the whole network is trainable."],
        ["Selection", "15% validation target, preferably complete recording groups. Best state is minimum validation cross entropy; stop after five stale epochs beyond warm up; gradient norm clipped at 1.0."],
        ["Reproducibility", "Python, NumPy and PyTorch seed 7; deterministic algorithm mode requested with warnings allowed."],
    ], [1.42, 5.62])
    equation(doc, "LCE = −(1/B) ∑ᵢ log p(yᵢ | Xᵢ)    θ* = argminₑ Lval(e)", "(6) Local training loss and best epoch selection")
    para(doc, "Validation detail for disclosure: the code computes alignment on the complete selected training pool before splitting train and validation recordings. No validation labels are used for this alignment, but validation EEG affects the unsupervised reference. If a recording level split lacks both classes, the code falls back to stratified trial level validation. The saved validation losses are therefore internal selection metrics, not a strict independent subject test.")

    heading(doc, "Exact Checkpoints Used in the Four Live Tests")
    table(doc, ["Live subject", "Training subjects and recordings", "Training and selection", "Checkpoint"], [
        ["S10", "S1, S3, S4, S5, S6, S7, S8; 28 recordings (runs 1–4 each). No S10 training labels.", "Best epoch 9; 14 epochs run. Validation accuracy 92.9%, loss 0.1780; fit 202.3 s.", "mirepnet__7subjects_28runs_22b10fee.pt"],
        ["S6", "S1, S3, S4, S5, S7, S8 runs 1–4 plus S10 runs 1–8; 32 recordings, 1,723 retained trials. No S6 training labels.", "Best epoch 10; 15 epochs run. Validation accuracy 86.1%, loss 0.2999; fit 237.9 s.", "mirepnet__finetuned_20260926_165237.pt"],
    ], [0.72, 2.48, 2.04, 1.80])
    para(doc, "Both checkpoints started independently from the same upstream pretrained weights. The S6 checkpoint was not a continuation of the S10 checkpoint. The S6 training manifest includes S10 runs 1–4; config.py and README.md mark those runs invalid for the frozen Local Exp4 validation cohort. Consequently the S6 live result is an observed result with that exact checkpoint, but it should not be described as a validated cohort result without a retraining analysis using only eligible runs.")
    para(doc, "Checkpoint integrity identifiers: S10 SHA 256 8f2da843a1ae57607c338673b4d67d9d427fba7f0f25a1ecd309e8f458b895e3; S6 SHA 256 46daeb62523ac88eb1b15e0b0ebbdb43f0a09d58989ceb6b6d272458fad3f596. Saved checkpoint metadata also contains classes, channel order, interpolation matrix, alignment reference, best epoch, loss, training time, recording list and history.", small=True)

    heading(doc, "Live Calibration and Command Generation")
    para(doc, "Each of the four sessions selected unlabeled calibration. The operator supplied 45 s of EEG and the code formed 84 overlapping windows, separated by 0.5 s. These windows fit a new 15 channel Euclidean Alignment reference for that session. Their logits also seeded a neutral decision center at the median log odds. The decoder weights were not changed by this calibration, and the calibration did not read left or right labels. These tests are therefore unseen subject with unlabeled adaptation, not strict no adaptation zero shot runs.")
    para(doc, "During control, the board supplies 125 Hz EEG. At approximately 0.2 s intervals, a 4 s live buffer is filtered at 8–30 Hz and resampled to 250 Hz. In task repeat mode the latest 2 s is duplicated to 1,000 samples, demeaned, aligned, interpolated and sent to the network. The buffer includes the extra earlier signal needed by the filter. The system logs raw EEG and every per window score.")
    equation(doc, "sₜ = log(p₂/p₁)    zₜ = sₜ − cₜ    q₂ = 1 / (1 + exp(−zₜ))", "(7) Log odds, adaptive neutral center and recentered decision score")
    para(doc, "The center starts at the calibration median. Only windows with |sₜ−cₜ| < log(0.65/0.35) update it by cₜ₊₁ = cₜ + 0.01(sₜ−cₜ); it is clamped within ±2 log odds of its seed. The recentered q values are used for control confidence. They are transformed decision scores and should not be reported as independently calibrated probabilities.")
    para(doc, "A window votes only when max(q₁,q₂) ≥0.85, equivalent to |z|≥log(0.85/0.15)≈1.735. A class requires at least four votes among the last five eligible or abstaining windows, then three consecutive consensus windows before commitment. The maze ignores votes computed while moving and requires a fully fresh post stop vote buffer before accepting one turn. A wrong turn is scored once; later corrective movements are not new model trials.")

    heading(doc, "Observed Live Maze Outcomes")
    scores = session_scores()
    rows = []
    for s in scores:
        rows.append([f"S{s['subject']} T{s['test']}", f"{s['maze']} / {s['seed']}",
                     f"{s['correct']}/40", percent(s['correct'], 40),
                     str(s['recovery']), f"{int(round(s['duration'])) // 60}:{int(round(s['duration'])) % 60:02d}"])
    table(doc, ["Session", "Maze / seed", "Correct", "Accuracy", "Recovery moves", "Wall time"], rows,
          [0.80, 1.10, 0.85, 0.90, 1.55, 0.90])
    para(doc, "S10 scored 66/80 (82.5%); S6 scored 77/80 (96.3% rounded); the combined descriptive count is 143/160 (89.4% rounded). Each test used a different standard maze: S10 mazes 1 and 2 (seeds 11 and 12); S6 mazes 3 and 4 (seeds 13 and 14). These are four sessions from two people, so the corner percentages are not independent participant level replication or a statistical comparison between checkpoints.")
    para(doc, "For S10, legacy logs included 20 and 18 subject correction turns after wrong model choices. For S6, the newer automatic correction produced 2 and 4 recovery turns. The corrected score counts only the first decision at each numbered corner. The face camera recordings are original; the maze videos in the accompanying report are deterministic replays of logged decisions and wait times because the original maze captures were corrupted.")

    heading(doc, "Reproducibility Notes and Limits")
    bullet(doc, "The official released pretrained encoder is external. Its upstream training data, subject screening and exact pretraining run were not generated by this project. Cite Liu et al. for that stage [1].")
    bullet(doc, "The 45 channel interpolation is a model compatibility transform. It cannot replace missing high quality measured EEG channels or recover unobserved spatial information.")
    bullet(doc, "The local training validation uses unlabeled EEG from held out recordings when fitting alignment; do not label its validation accuracy as strict independent or prospective.")
    bullet(doc, "During the first two head only epochs, encoder parameter gradients are disabled, but the module remains in training mode. Batch normalization running statistics in the embedding can still update.")
    bullet(doc, "Offline trials exceeding 100 μV peak to peak are rejected. Live inference does not apply that rejection rule to each window, so electrode quality remains an operational dependency.")
    bullet(doc, "Live calibration and adaptive decision recentering use the target subject's EEG without labels. State this when describing the unseen subject protocol.")
    bullet(doc, "The two live subjects used different checkpoint training pools and different maze layouts. Their 82.5% and 96.3% scores are descriptive, not an isolated estimate of person or model effects.")
    bullet(doc, "The S6 checkpoint includes S10 runs 1–4 even though the repository's frozen data contract excludes them. Preserve this provenance in any paper table, or retrain an eligible checkpoint and repeat the evaluation.")
    para(doc, "Core implementation and evidence: classifier/mirepnet.py; classifier/run.py; classifier/smoother.py; config.py; scripts/finetune_mirepnet_selected.py; simulation/src/tiago_maze/ws_client.py; simulation/src/tiago_maze/scoring.py; models/mirepnet/*.pt and *.training.json; paper results/subject_10_test_{1,2} and subject_6_test_{1,2}/results summary.json. Dependency versions are pinned in pyproject.toml and uv.lock. The current working tree contains uncommitted changes, so archive a commit or code snapshot with the paper.", small=True)

    heading(doc, "References")
    para(doc, "[1] D. Liu et al., “MIRepNet: A pipeline and pre-trained model for EEG-based motor imagery classification,” Knowledge-Based Systems, vol. 343, art. 115966, 2026. doi: 10.1016/j.knosys.2026.115966.", small=True)
    para(doc, "[2] Braindecode, “MIRepNet pretrained model card,” Hugging Face, revision 857f1e3642976be9f2ca6883e5508c3f7c91d86f. https://huggingface.co/braindecode/mirepnet-pretrained", small=True)

    doc.core_properties.title = "MIRepNet Model Implementation and Methods"
    doc.core_properties.subject = "IEEE manuscript methods and implementation disclosure"
    doc.core_properties.keywords = "MIRepNet, motor imagery, EEG, Euclidean Alignment, transformer, OpenBCI"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
